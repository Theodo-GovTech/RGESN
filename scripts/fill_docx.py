"""Génère la déclaration d'écoconception publique (docx) depuis un declaration.json.

Usage:
    python scripts/fill_docx.py <declaration.json> [<output.docx>]

Le document produit suit le plan de docs/rgesn_2024_exemple_declaration.docx :
  - Titre + métadonnées service
  - Résumé (critères validés / non validés)
  - Détails du diagnostic regroupés par thème (1 Stratégie → 9 Algorithmie)

Si <output.docx> n'est pas fourni, le fichier est écrit dans out/<service>_<date>.docx.
"""
from __future__ import annotations

import json
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
CRITERES = ROOT / "data" / "criteres_rgesn.json"

VALID_EVAL = {"Conforme", "Non conforme", "Non applicable", "À évaluer"}


def slugify(s: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", s).strip("_") or "service"


def load_declaration(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if "service" not in data or "criteres" not in data:
        raise ValueError("declaration.json invalide : 'service' et 'criteres' requis")
    for c in data["criteres"]:
        if c.get("evaluation") and c["evaluation"] not in VALID_EVAL:
            raise ValueError(f"Évaluation invalide pour {c.get('id')} : {c['evaluation']!r}")
    return data


def load_criteres_index() -> OrderedDict:
    raw = json.loads(CRITERES.read_text(encoding="utf-8"))
    index = OrderedDict()
    for c in raw:
        index[c["id"]] = c
    return index


def id_sort_key(cid: str) -> tuple:
    a, b = cid.split(".")
    return int(a), int(b)


def add_meta_line(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{label} : ")
    run.bold = True
    p.add_run(value or "—")


def add_critere(doc: Document, meta: dict, entry: dict, default_date: str) -> None:
    doc.add_heading(meta["libelle"], level=3)

    marker = doc.add_paragraph()
    evaluation = entry.get("evaluation", "À évaluer")
    eval_date = entry.get("date_evaluation") or default_date
    run = marker.add_run(f"Critère {meta['id']} — {evaluation} ({eval_date})")
    run.italic = True
    run.font.size = Pt(9)

    texte = entry.get("texte_declaration") or meta.get("exemple_declaration") or "[À compléter]"
    for para in texte.split("\n"):
        doc.add_paragraph(para)

    if entry.get("evolutions_potentielles"):
        p = doc.add_paragraph()
        r = p.add_run("Évolutions potentielles : ")
        r.bold = True
        p.add_run(entry["evolutions_potentielles"])


def build(declaration: dict, index: OrderedDict) -> Document:
    service = declaration["service"]
    eval_date = service.get("date_evaluation") or date.today().isoformat()

    entries = {c["id"]: c for c in declaration["criteres"]}

    doc = Document()

    doc.add_heading(f"Déclaration d'écoconception de {service.get('nom', '[service]')}", level=0)
    doc.add_paragraph(f"Date de réalisation : {eval_date}")

    doc.add_heading("Résumé", level=1)

    doc.add_heading("Objectif", level=2)
    doc.add_paragraph(
        f"Le service {service.get('nom', '[service]')} s'inscrit dans une démarche d'écoconception "
        f"visant à réduire ses impacts environnementaux. Cette déclaration a été rédigée le "
        f"{eval_date}, dans le cadre de la mise en œuvre du référentiel général d'écoconception "
        f"des services numériques (RGESN, version 2024, Arcep/Arcom)."
    )

    doc.add_heading("Chemins critiques et unités fonctionnelles évalués", level=2)
    add_meta_line(doc, "Échantillons", service.get("echantillons", ""))
    add_meta_line(doc, "Entité qui évalue", service.get("entite_evaluation", ""))
    add_meta_line(doc, "Responsable de l'évaluation", service.get("responsable_evaluation", ""))

    valides = sorted(
        (cid for cid, c in entries.items() if c.get("evaluation") == "Conforme"),
        key=id_sort_key,
    )
    non_valides = sorted(
        (cid for cid, c in entries.items() if c.get("evaluation") == "Non conforme"),
        key=id_sort_key,
    )
    non_applicables = sorted(
        (cid for cid, c in entries.items() if c.get("evaluation") == "Non applicable"),
        key=id_sort_key,
    )
    a_evaluer = sorted(
        (cid for cid, c in entries.items() if c.get("evaluation") == "À évaluer"),
        key=id_sort_key,
    )

    doc.add_heading("Critères validés par le service numérique", level=2)
    doc.add_paragraph(", ".join(valides) if valides else "—")

    doc.add_heading("Critères non validés par le service numérique", level=2)
    doc.add_paragraph(", ".join(non_valides) if non_valides else "—")

    if non_applicables:
        doc.add_heading("Critères non applicables", level=2)
        doc.add_paragraph(", ".join(non_applicables))

    if a_evaluer:
        doc.add_heading("Critères restant à évaluer", level=2)
        doc.add_paragraph(", ".join(a_evaluer))

    doc.add_heading(
        "Détails du diagnostic avec le référentiel général de l'écoconception des services numériques",
        level=1,
    )

    by_theme: dict[str, list[str]] = defaultdict(list)
    for cid in entries:
        meta = index.get(cid)
        if meta:
            by_theme[meta["theme"]].append(cid)

    for theme in sorted(by_theme, key=lambda t: int(t.split(" ", 1)[0])):
        doc.add_heading(theme, level=2)
        for cid in sorted(by_theme[theme], key=id_sort_key):
            add_critere(doc, index[cid], entries[cid], eval_date)

    return doc


def fill(declaration_path: Path, output_path: Path | None = None) -> Path:
    declaration = load_declaration(declaration_path)
    index = load_criteres_index()

    service = declaration["service"]
    eval_date = service.get("date_evaluation") or date.today().isoformat()

    if output_path is None:
        slug = slugify(service.get("nom", "service"))
        output_path = ROOT / "out" / f"{slug}_{eval_date}.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = build(declaration, index)
    doc.save(output_path)

    unknown = [c["id"] for c in declaration["criteres"] if c["id"] not in index]
    written = len(declaration["criteres"]) - len(unknown)
    rel = output_path.relative_to(ROOT) if output_path.is_relative_to(ROOT) else output_path
    print(f"OK {written} critères écrits dans {rel}")
    if unknown:
        print(f"  Avertissement : ids inconnus ignorés : {', '.join(unknown)}")
    return output_path


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    decl = Path(sys.argv[1]).resolve()
    out = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else None
    fill(decl, out)


if __name__ == "__main__":
    main()
