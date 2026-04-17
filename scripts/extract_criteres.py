"""Extrait les 78 critères du tableur officiel RGESN 2024 vers data/criteres_rgesn.json.

Source: docs/rgesn_2024_outil_declaration.xlsx (Arcep, mai 2024).
Chaque critère occupe 2 lignes dans le xlsx :
  - ligne N   : B=id, C=libellé, D=priorité, E=évaluation (vide), F=date (vide),
                G=évolutions (vide), H=difficulté, M=pondération, N=cible, O=moyen de test
  - ligne N+1 : C:G fusionnée (texte de déclaration, à remplir) ; H:N fusionnée (exemple de déclaration)
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "rgesn_2024_outil_declaration.xlsx"
OUT = ROOT / "data" / "criteres_rgesn.json"


def clean_theme(sheet_name: str) -> str:
    return re.sub(r"\s*\(\d+\)\s*$", "", sheet_name).strip()


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    criteria = []
    for sn in wb.sheetnames:
        if sn == "Score d'avancement":
            continue
        ws = wb[sn]
        for row_idx in range(1, ws.max_row + 1):
            b = ws.cell(row=row_idx, column=2).value
            if not (isinstance(b, str) and re.match(r"^\d+\.\d+$", b.strip())):
                continue
            cont = row_idx + 1
            exemple = ws.cell(row=cont, column=8).value
            criteria.append({
                "id": b.strip(),
                "theme": clean_theme(sn),
                "feuille_xlsx": sn,
                "ligne_xlsx": row_idx,
                "libelle": (ws.cell(row=row_idx, column=3).value or "").strip(),
                "priorite": (ws.cell(row=row_idx, column=4).value or "").strip(),
                "difficulte": (ws.cell(row=row_idx, column=8).value or "").strip(),
                "ponderation": ws.cell(row=row_idx, column=13).value,
                "cible": (ws.cell(row=row_idx, column=14).value or "").strip(),
                "moyen_test": (ws.cell(row=row_idx, column=15).value or "").strip(),
                "exemple_declaration": exemple.strip() if isinstance(exemple, str) else "",
            })
    criteria.sort(key=lambda c: tuple(int(x) for x in c["id"].split(".")))
    OUT.write_text(json.dumps(criteria, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(criteria)} critères écrits dans {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
